"""Build editable PPTX, vector PDF and previews from one slide layout.

Dependencies: python-pptx, reportlab, pillow. No API calls or generated data.
Run from repo root: PYTHONPATH=/private/tmp/crypto-slides-deps .venv/bin/python docs/presentations/b1_b2_2026_09_21/build_slides.py
"""
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

OUT = Path(__file__).resolve().parent
FONT = "/System/Library/Fonts/Supplemental/Arial.ttf"
BOLD = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
pdfmetrics.registerFont(TTFont("ArialLocal", FONT))
pdfmetrics.registerFont(TTFont("ArialLocalBold", BOLD))
W, H = 960, 540
C = {"bg": "F5F6FA", "white": "FFFFFF", "ink": "1D2433", "muted": "606B7E",
     "red": "9F2239", "teal": "147D79", "amber": "AC6B12", "line": "DDE2EA",
     "pink": "F8E9ED", "green": "E7F3F0", "gold": "FFF3DF", "blue": "EAF0FC"}
slides = []


def new(title, section, source="", notes=""):
    slide = {"title": title, "notes": notes, "elements": []}
    slides.append(slide)
    rect(0, 0, W, H, C["bg"])
    rect(0, 0, 8, H, C["red"])
    text(42, 23, 860, section.upper(), 10, C["red"], True)
    text(42, 49, 880, title, 29, bold=True)
    rect(42, 492, 876, 1, C["line"])
    text(42, 504, 760, source or "Trần Chí Vũ · 20236060 · Multi-Agent Crypto", 9, C["muted"])
    text(870, 502, 48, f"{len(slides):02d} / 12", 10, C["muted"], True)
    return slide


def rect(x, y, w, h, color):
    slides[-1]["elements"].append({"type": "rect", "x": x, "y": y, "w": w, "h": h, "color": color})


def text(x, y, w, value, size=19, color=None, bold=False):
    font = "ArialLocalBold" if bold else "ArialLocal"
    lines = []
    for paragraph in value.split("\n"):
        line = ""
        for word in paragraph.split():
            attempt = f"{line} {word}".strip()
            if line and pdfmetrics.stringWidth(attempt, font, size) > w-6:
                lines.append(line)
                line = word
            else:
                line = attempt
        lines.append(line)
    height = len(lines)*size*1.22+5
    if y+height > H:
        raise ValueError(f"Text outside slide {len(slides)}: {value}")
    slides[-1]["elements"].append({"type": "text", "x": x, "y": y, "w": w, "h": height,
                                 "lines": lines, "size": size, "bold": bold, "color": color or C["ink"]})
    return height


def card(x, y, w, h, label, body, color=None, body_size=18):
    rect(x, y, w, h, color or C["white"])
    heading_h = text(x+18, y+16, w-36, label, 20, bold=True)
    body_h = text(x+18, y+23+heading_h, w-36, body, body_size, C["muted"])
    if 23+heading_h+body_h > h-10:
        raise ValueError(f"Card overflow on slide {len(slides)}: {label}")


# 01 — cover
new("", "BỘ QUY TẮC ĐÁNH GIÁ · BƯỚC B1", notes="Em trình bày khung đánh giá hệ thống Multi-Agent Crypto: các nhóm metric, quy tắc xử lý tín hiệu, mô phỏng giao dịch và thiết kế so sánh. Mục tiêu là thống nhất cách đo trước khi đánh giá hiệu quả giải pháp.")
rect(8, 0, 952, 492, C["ink"])
text(48, 35, 850, "BỘ QUY TẮC ĐÁNH GIÁ  /  BƯỚC B1", 12, "E7A8B6", True)
text(48, 93, 850, "Đánh giá hiệu quả\nMulti-Agent Crypto", 42, C["white"], True)
text(50, 219, 760, "Bộ metric · Quy tắc mô phỏng · Thiết kế đánh giá", 23, "CBD3DF")
text(50, 276, 760, "Trần Chí Vũ  |  MSSV: 20236060", 19, C["white"])
for x, label, body in [(48, "DỰ ĐOÁN", "Đúng hướng giá?"), (348, "TÀI CHÍNH", "Có lãi sau chi phí?"), (648, "CƠ CHẾ", "Debate đóng góp gì?")]:
    rect(x, 357, 268, 89, "2C3547")
    text(x+16, 372, 238, label, 12, "E7A8B6", True)
    text(x+16, 398, 238, body, 18, C["white"])

# 02 — questions / system flow
new("Ba câu hỏi cần trả lời", "01 / Mục tiêu đánh giá", notes="Hệ thống có ba specialist cho on-chain, thị trường và tâm lý. Validator đo bất đồng và chỉ kích hoạt Debate khi vượt ngưỡng. Mục tiêu đánh giá tách thành dự báo, giao dịch mô phỏng và đóng góp của Debate. Không suy ra lợi nhuận từ accuracy.")
card(42, 112, 280, 164, "01  Dự đoán", "BUY/SELL có đúng hướng giá sau 24 giờ không?", C["white"], 21)
card(340, 112, 280, 164, "02  Giao dịch", "Theo tín hiệu có lãi sau chi phí mô phỏng không?", C["white"], 21)
card(638, 112, 280, 164, "03  Debate", "Tranh biện có cải thiện quyết định cuối không?", C["white"], 21)
for x,w,label,fill in [(42,228,"3 specialist",C["blue"]),(301,172,"Validator",C["blue"]),(504,172,"Debate*",C["pink"]),(707,211,"Mediator",C["green"])]:
    rect(x, 322, w, 58, fill)
    text(x+16, 338, w-26, label, 21, bold=True)
for x in [277,480,684]:
    text(x, 337, 25, "→", 21)
text(42, 402, 875, "BTC · 1 quyết định/ngày · 00:00 UTC · Chân trời dự báo 24 giờ", 18, C["ink"], True)
text(42, 443, 875, "* Chỉ chạy Debate khi conflict vượt ngưỡng; nếu không, chuyển thẳng sang Mediator.", 15, C["muted"])

# 03 — metrics
new("Bộ metric gồm ba nhóm", "02 / Bộ chỉ số B1", notes="Directional accuracy là chỉ số chính của dự đoán. Nhóm tài chính đo hiệu quả giao dịch mô phỏng theo quy tắc cố định. S_final và conflict score giúp giải thích cơ chế, không phải bằng chứng chất lượng tự thân. Tỷ lệ đổi tín hiệu phải đi cùng kết quả sai sang đúng và đúng sang sai.")
card(42, 116, 280, 310, "CHẤT LƯỢNG DỰ ĐOÁN", "Directional accuracy\n\nBUY / SELL accuracy\n\nDirectional coverage", C["white"], 19)
card(340, 116, 280, 310, "TÀI CHÍNH MÔ PHỎNG", "Cumulative return\nTrade win rate\nProfit factor\nMaximum drawdown\n\nTính sau phí", C["white"], 19)
card(638, 116, 280, 310, "CƠ CHẾ HỆ THỐNG", "S_final · Conflict score\nTỷ lệ đổi tín hiệu\nSai → đúng / đúng → sai\nChuyển đổi NEUTRAL\nTỷ lệ chạy thành công", C["white"], 18)
text(42, 448, 876, "Đúng hướng ≠ lệnh có lãi. Conflict giảm ≠ dự đoán tốt hơn.", 20, C["red"], True)

# 04 — handling states
new("Chấm đúng mẫu số, giữ nguyên ngày lỗi", "03 / Quy tắc dữ liệu", notes="BUY đúng khi lợi suất tương lai dương, SELL đúng khi âm. Giá không đổi được tính là sai với BUY/SELL. NEUTRAL và lỗi không nằm trong mẫu số directional accuracy. Coverage tính trên toàn bộ ngày yêu cầu, chỉ tính BUY/SELL hợp lệ vào tử số. Tỷ lệ thành công có cả NEUTRAL. Accuracy cặp dùng các ngày hợp lệ chung; metric tài chính phải giữ nguyên lịch, không nối tắt ngày lỗi. Chỉ số không có mẫu số được ghi N/A hoặc null.")
xs=[42,205,395,565,727]; ws=[163,190,170,162,191]
headers=["TRẠNG THÁI","ACCURACY","COVERAGE","THÀNH CÔNG","GIAO DỊCH"]
for x,w,h in zip(xs,ws,headers):
    rect(x,116,w,42,C["ink"]);text(x+12,128,w-20,h,12,C["white"],True)
rows=[("BUY / SELL","Chấm đúng/sai","Có","Có","Mở / giữ lệnh"),
      ("NEUTRAL","Không chấm","Không","Có","Đóng → nghỉ"),
      ("Lỗi","Không chấm","Không","Không","Đóng → nghỉ")]
for i,row in enumerate(rows):
    for x,w,v in zip(xs,ws,row):
        rect(x,160+i*53,w-2,51,C["white"] if i!=2 else C["pink"])
        text(x+12,175+i*53,w-20,v,17,bold=(x==42))
card(42, 346, 428, 124, "Coverage", "Ngày BUY/SELL hợp lệ\n÷ toàn bộ ngày yêu cầu", C["green"], 18)
card(490, 346, 428, 124, "Tỷ lệ chạy thành công", "Ngày hợp lệ, kể cả NEUTRAL\n÷ toàn bộ ngày yêu cầu", C["blue"], 18)

# 05 — trading
new("Quy tắc giao dịch được cố định trước", "04 / Mô phỏng tài chính", notes="BUY là long, SELL là short. Một lệnh là chuỗi ngày cùng hướng, số lượng BTC giữ nguyên. NEUTRAL hoặc ngày lỗi đóng lệnh. Quy mô chừa vốn trả phí mở: q bằng E chia P mở nhân một cộng f. Phí trên giá trị mỗi lần giao dịch, đóng cuối kỳ. Short không được nhân dồn một trừ return ngày. Khi cháy vốn, cap equity ở 0, lưu riêng PnL công thức và ghi nhận. Đây là mô phỏng lý tưởng, không phải giao dịch thực tế.")
card(42, 114, 426, 203, "Vị thế và thời gian", "BUY = long · SELL = short\nCùng tín hiệu → giữ nguyên q\nNEUTRAL / lỗi → đóng lệnh\nĐóng mọi vị thế cuối kỳ", C["white"], 18)
card(490, 114, 428, 203, "Vốn và chi phí", "q = E / [P_mở × (1 + f)]\nVốn tích lũy qua các lệnh\nPhí: 0% · 0,1% · 0,2% mỗi chiều\nMức cơ sở: 0,1%", C["white"], 18)
rect(42, 337, 876, 67, C["green"])
text(60, 350, 840, "PnL = q × biến động giá theo hướng vị thế − phí mở − phí đóng", 22, C["teal"], True)
text(42, 424, 876, "Giả định: độ trễ bằng 0; bỏ trượt giá, funding và chi phí vay short.\nDrawdown đo theo ngày, gồm lãi/lỗ chưa thực hiện và phí đã trả.", 17, C["muted"])

# 06 — prediction horizon and ground truth
new("Một quyết định, một nhãn sau 24 giờ", "05 / Ground truth", notes="Mỗi ngày có một thời điểm quyết định lúc 00:00 UTC. Đầu vào chỉ dùng thông tin đã sẵn có trước thời điểm quyết định; cần xét thời gian công bố, không chỉ ngày ghi trên dữ liệu. P_t là giá đóng cửa ngày trước dùng làm mốc mô phỏng, P_t+24h là giá tại mốc kế tiếp. Lợi suất tương lai chỉ dùng để chấm điểm, không đưa vào đầu vào agent. Cách vào lệnh tại giá mốc là giả định độ trễ bằng không.")
card(42, 115, 280, 214, "TRƯỚC t", "Giá và chỉ báo lịch sử\nOn-chain / tâm lý\nKiểm tra thời gian công bố", C["blue"], 19)
card(340, 115, 280, 214, "TẠI t · 00:00 UTC", "1 tín hiệu mỗi ngày\nBUY / SELL / NEUTRAL\nLưu riêng trạng thái lỗi", C["white"], 19)
card(638, 115, 280, 214, "TẠI t + 24 GIỜ", "Quan sát giá tương lai\nTính lợi suất thực tế\nĐối chiếu hướng dự đoán", C["green"], 19)
rect(42, 355, 876, 63, C["ink"])
text(62, 369, 834, "r_24h = P_(t+24h) / P_t − 1", 26, C["white"], True)
text(42, 443, 875, "Dữ liệu tương lai chỉ dùng làm nhãn đánh giá, không đưa vào đầu vào agent.", 18, C["red"], True)

# 07 — prediction scoring
new("Chấm hướng và báo độ bao phủ cùng nhau", "06 / Chất lượng dự đoán", notes="Directional accuracy bằng số dự đoán đúng chia số ngày hợp lệ có BUY hoặc SELL. BUY đúng nếu lợi suất dương, SELL đúng nếu lợi suất âm; lợi suất bằng không được tính là sai. BUY accuracy và SELL accuracy được báo riêng để phát hiện lệch hướng. Nếu không có dự đoán phù hợp, chỉ số tương ứng là N/A, lưu JSON là null. Luôn báo số đúng, tổng số dự đoán, coverage và tỷ lệ thành công để tránh diễn giải accuracy thiếu mẫu số.")
rect(42, 114, 876, 83, C["ink"])
text(62, 132, 834, "Directional accuracy = số dự đoán đúng / số BUY–SELL hợp lệ", 23, C["white"], True)
card(42, 219, 280, 179, "BUY accuracy", "Số BUY gặp r > 0\n÷ số BUY hợp lệ", C["white"], 21)
card(340, 219, 280, 179, "SELL accuracy", "Số SELL gặp r < 0\n÷ số SELL hợp lệ", C["white"], 21)
card(638, 219, 280, 179, "Trường hợp biên", "r = 0 → chấm sai\nMẫu số = 0 → N/A", C["gold"], 20)
text(42, 432, 876, "Báo kèm: số đúng / số dự đoán · Coverage · Tỷ lệ chạy thành công", 19, C["red"], True)

# 08 — controlled comparison
new("Bốn cấu hình để so sánh phương pháp", "07 / Baseline và đối chứng", notes="Bốn cấu hình gồm luôn BUY, luôn SELL, nhánh không Debate và nhánh đầy đủ. Hai nhánh đa agent dùng chung outputs specialist. No-debate vẫn tính conflict nhưng bỏ RCA LLM. Để so chất lượng dự báo theo cặp, dùng các ngày mà hai nhánh có đầu ra hợp lệ và báo riêng chuyển đổi NEUTRAL. Coverage, tỷ lệ thành công và metric tài chính vẫn dùng toàn bộ lịch; không bỏ ngày lỗi rồi nối các ngày còn lại.")
for x,y,label,body in [(42,113,"Luôn BUY","Long xuyên kỳ theo quy tắc phí"),(490,113,"Luôn SELL","Short xuyên kỳ theo quy tắc phí"),(42,233,"Không Debate","Specialist → Validator → Mediator"),(490,233,"Có Debate","Cùng specialist + Debate có điều kiện")]:
    card(x,y,428,103,label,body,C["white"],16)
rect(42,365,876,98,C["blue"])
text(60,378,840,"Hai nhánh đa agent dùng chung đầu ra specialist",23,C["ink"],True)
text(60,418,840,"So accuracy theo cặp trên ngày hợp lệ chung.\nCoverage, tỷ lệ thành công và tài chính giữ đủ lịch.",17,C["muted"])

# 09 — Debate assessment
new("Tách thay đổi quyết định khỏi cải thiện chất lượng", "08 / Đóng góp của Debate", notes="Tỷ lệ đổi tín hiệu là số ngày tín hiệu sau khác trước Debate chia số ngày có Debate hợp lệ. Nó đo mức độ tác động, chưa chứng minh cải thiện. Trên các cặp BUY/SELL có thể chấm ở cả hai nhánh, báo số sai sang đúng và đúng sang sai. Tách riêng chuyển đổi liên quan NEUTRAL và báo coverage trên toàn bộ lịch. S_final được phân tích theo độ lớn và từng chiều; conflict mô tả mức bất đồng, không thay cho accuracy.")
card(42,114,428,158,"Tác động đến quyết định","Tỷ lệ đổi tín hiệu sau Debate\nS_final trước / sau\nConflict score ban đầu",C["white"],19)
card(490,114,428,158,"Tác động đến chất lượng","Sai → đúng và đúng → sai\nDirectional accuracy theo cặp\nChuyển đổi NEUTRAL + coverage",C["white"],19)
rect(42,295,876,87,C["blue"])
text(60,310,840,"Tỷ lệ đổi = số ngày đổi tín hiệu / số ngày có Debate hợp lệ",22,C["ink"],True)
text(42,411,876,"Đổi nhiều chưa chắc tốt hơn. Cần đối chiếu với nhãn giá tương lai.\nPhân tích |S_final| riêng cho chiều BUY và SELL.",19,C["red"],True)

# 10 — test case design
new("Test case kiểm tra cả phép tính và phương pháp", "09 / Thiết kế kiểm thử", notes="Test phép tính dùng chuỗi giá và tín hiệu nhỏ để tính tay được. Test đánh giá phương pháp dùng dữ liệu thị trường thật và bao phủ tăng, giảm, đi ngang. Hai loại có mục đích khác nhau: ví dụ tính tay kiểm tra công thức, còn dữ liệu lịch sử đánh giá chất lượng dự đoán. Khoảng 90 ngày, cách chia trạng thái thị trường, prompt, ngưỡng và phí phải được chốt trước khi chạy đánh giá.")
card(42,114,428,243,"KIỂM TRA PHÉP TÍNH","BUY / SELL gặp tăng, giảm, đi ngang\nNEUTRAL, lỗi, không có lệnh\nGiữ lệnh, đảo chiều, đóng cuối kỳ\nPhí, hòa vốn, không có lệnh lỗ\nShort, drawdown, vốn về 0",C["white"],18)
card(490,114,428,243,"ĐÁNH GIÁ PHƯƠNG PHÁP","Thị trường tăng / giảm / đi ngang\nLuôn BUY / luôn SELL\nCó Debate / không Debate\nĐộ nhạy phí: 0% / 0,1% / 0,2%\nCùng quy tắc và lịch đánh giá",C["white"],18)
rect(42,384,876,80,C["ink"])
text(61,398,835,"Chốt trước: khoảng 90 ngày và quy tắc chọn giai đoạn,",22,C["white"],True)
text(61,430,835,"cùng prompt, ngưỡng tín hiệu và giả định giao dịch.",20,"CBD3DF")

# 11 — appendix calculation
new("Phụ lục · Ba ví dụ tính tay", "A / Cơ sở kiểm thử", notes="Các số liệu ở đây là ví dụ kiểm tra công thức, không phải kết quả hệ thống. Short giữ cố định q ở phí zero có return 5% khi giá từ 100 đến 95, bất kể qua 90. Test phí cần q chừa phí đầu vào. Profit factor phải tổng hợp PnL tiền khi vốn tích lũy, không cộng tỷ lệ phần trăm từng lệnh.")
card(42,115,280,292,"VÍ DỤ 1 · SHORT","Giá: 100 → 90 → 95\nPhí 0%, giữ nguyên q\n\nLợi suất = 1 − 95/100\n\nĐáp án: +5%",C["white"],18)
card(340,115,280,292,"VÍ DỤ 2 · PHÍ","Long: 100 → 110\nE = 1; f = 0,1%\nq = 1 / (100 × 1,001)\n\nPnL ≈ 0,0978022\nĐáp án: +9,78022%",C["white"],18)
card(638,115,280,292,"VÍ DỤ 3 · PF","Lợi suất ròng các lệnh:\n+10%, −4%, +6%, −6%\nVốn tích lũy, E đầu = 1\n\nPF = 0,16336 / 0,1111616\nĐáp án: 1,46957",C["white"],17)
text(42,435,875,"Còn kiểm tra: NEUTRAL · Ngày lỗi · Đảo chiều · Hòa vốn · Cháy vốn",18,C["muted"])

# 12 — reporting conventions
new("Phụ lục · Quy ước báo cáo kết quả", "B / Cách diễn giải metric", notes="Accuracy cần đi kèm số đúng và mẫu số, coverage và tỷ lệ thành công. Metric tài chính dùng PnL ròng theo tiền sau phí, với drawdown từ đường vốn hằng ngày gồm lãi lỗ chưa thực hiện. Nếu chỉ có lệnh thắng mà không có lệnh lỗ, profit factor hiển thị vô cùng nhưng JSON lưu null kèm lý do. Không có lệnh hoặc toàn hòa vốn thì PF là N/A. Đầu vào, cấu hình và đầu ra từng ngày cần được lưu để đối chiếu kết quả.")
card(42,114,876,159,"Báo đủ thông tin để đối chiếu","Accuracy: số đúng / số dự đoán + coverage + tỷ lệ thành công\nTài chính: PnL ròng, số lệnh, phí và đường vốn theo ngày\nDebate: cặp trước / sau, sai ↔ đúng, chuyển đổi NEUTRAL",C["white"],18)
card(42,291,876,181,"Trường hợp không xác định","Mẫu số bằng 0 → N/A; lưu JSON là null kèm lý do.\nPF: có lãi, không có lỗ → hiển thị ∞; JSON lưu null + lý do.\nKhông có lệnh hoặc toàn hòa vốn → PF là N/A.\nLưu dữ liệu đầu vào, cấu hình và đầu ra từng ngày để truy vết.",C["gold"],18)


def rgb(value):
    return tuple(int(value[i:i+2],16) for i in (0,2,4))


def export():
    OUT.mkdir(parents=True,exist_ok=True)
    preview=OUT/"preview";preview.mkdir(exist_ok=True)
    prs=Presentation();prs.slide_width=Pt(W);prs.slide_height=Pt(H)
    prs.core_properties.title="Multi-Agent Crypto — Bộ quy tắc đánh giá B1"
    prs.core_properties.author="Trần Chí Vũ — 20236060"
    pdf=canvas.Canvas(str(OUT/"Bao_cao_B1_ke_hoach_ngay_3_4.pdf"),pagesize=(W,H))
    pdf.setTitle(prs.core_properties.title);pdf.setAuthor(prs.core_properties.author)
    notes=["# Ghi chú thuyết trình — Bộ quy tắc đánh giá B1", "", "10 slide chính: khoảng 7–10 phút. Slide 11–12 là phụ lục.", ""]
    for index, slide in enumerate(slides,1):
        ps=prs.slides.add_slide(prs.slide_layouts[6])
        ps.notes_slide.notes_text_frame.text=slide["notes"]
        im=Image.new("RGB",(W*2,H*2),"white");d=ImageDraw.Draw(im)
        for item in slide["elements"]:
            x,y,w,h=item["x"],item["y"],item["w"],item["h"]
            color=rgb(item["color"])
            if item["type"]=="rect":
                shape=ps.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h))
                shape.fill.solid();shape.fill.fore_color.rgb=RGBColor(*color);shape.line.fill.background()
                pdf.setFillColorRGB(*(v/255 for v in color));pdf.rect(x,H-y-h,w,h,fill=1,stroke=0)
                d.rectangle((x*2,y*2,(x+w)*2,(y+h)*2),fill=color)
            else:
                size=item["size"];font="ArialLocalBold" if item["bold"] else "ArialLocal"
                box=ps.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h));tf=box.text_frame
                tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.word_wrap=False
                for k,line in enumerate(item["lines"]):
                    paragraph=tf.paragraphs[0] if k==0 else tf.add_paragraph()
                    paragraph.text=line;paragraph.font.name="Arial";paragraph.font.size=Pt(size)
                    paragraph.font.bold=item["bold"];paragraph.font.color.rgb=RGBColor(*color)
                    paragraph.line_spacing=Pt(size*1.22);paragraph.space_after=Pt(0);paragraph.space_before=Pt(0)
                    ly=y+k*size*1.22
                    pdf.setFont(font,size);pdf.setFillColorRGB(*(v/255 for v in color));pdf.drawString(x,H-ly-size*.92,line)
                    pf=ImageFont.truetype(BOLD if item["bold"] else FONT,round(size*2))
                    d.text((x*2,(ly+size*.92)*2),line,font=pf,fill=color,anchor="ls")
        pdf.showPage();im.save(preview/f"slide_{index:02d}.png")
        notes.extend([f"## Slide {index:02d} — {slide['title'] or 'Đánh giá hiệu quả Multi-Agent Crypto'}","",slide["notes"],""])
    pdf.save();prs.save(OUT/"Bao_cao_B1_ke_hoach_ngay_3_4.pptx")
    (OUT/"GHI_CHU_THUYET_TRINH.md").write_text("\n".join(notes),encoding="utf-8")
    sheet=Image.new("RGB",(1440,4*290),"#DDE2EA")
    for i in range(12):
        im=Image.open(preview/f"slide_{i+1:02d}.png");im.thumbnail((464,261))
        sheet.paste(im,(8+(i%3)*480,8+(i//3)*290))
    sheet.save(preview/"contact_sheet.png")
    (OUT/"README.md").write_text("# Slide báo cáo bộ quy tắc đánh giá B1\n\n- `Bao_cao_B1_ke_hoach_ngay_3_4.pptx`: PowerPoint chỉnh sửa được, có speaker notes.\n- `Bao_cao_B1_ke_hoach_ngay_3_4.pdf`: bản trình chiếu PDF 16:9.\n- `GHI_CHU_THUYET_TRINH.md`: lời dẫn cho từng slide.\n- `preview/`: ảnh xem nhanh.\n\n10 slide chính + 2 phụ lục; thời lượng khoảng 7–10 phút. Nội dung tập trung vào bộ metric, ground truth, quy tắc mô phỏng, baseline, Debate và test case. Ví dụ tính tay chỉ minh họa công thức. Giữ tên tệp để các liên kết đã gửi vẫn hoạt động.\n",encoding="utf-8")
    print(f"Exported {len(slides)} slides to {OUT}")


if __name__=="__main__":
    export()
